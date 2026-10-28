import argparse
import json
import math
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


LABELS = {
    "ramas": "Ramas",
    "tags": "Tags",
    "referencias": "Referencias Git",
    "commits": "Commits",
    "objetos": "Objetos y contenido Git",
    "integridad_gitlab": "Integridad del origen GitLab",
    "integridad_github": "Integridad del destino GitHub",
    "rama_predeterminada": "Rama predeterminada",
    "variables_y_secrets": "Variables y Secrets (alcance del validador)",
    "environments": "Environments y politicas configuradas",
    "protecciones_de_rama": "Protecciones de rama (alcance del validador)",
    "miembros_y_permisos": "Miembros y permisos",
    "pipelines": "Pipelines y workflows",
    "lfs": "Git LFS",
    "disponibilidad_para_certificar": "Disponibilidad para verificar",
}
COLORS = {"OK": "E0F2E9", "ERROR": "FCE4E4", "ADVERTENCIA": "FFF1CC"}
CONTROL_CHARACTERS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def sheet_title(repository, used):
    base = re.sub(r"[\[\]:*?/\\]", "_", repository).strip("'")[:31] or "Repositorio"
    title = base
    sequence = 2
    while title.casefold() in used or title.casefold() == "history":
        suffix = f" ({sequence})"
        title = base[: 31 - len(suffix)] + suffix
        sequence += 1
    used.add(title.casefold())
    return title


def write_cell(sheet, row, column, value):
    cell = sheet.cell(row, column)
    if isinstance(value, str):
        cell.value = CONTROL_CHARACTERS.sub("", value)[:30000]
        cell.data_type = "s"
    else:
        cell.value = value
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    return cell


def append_row(sheet, values, fill=None):
    row = sheet.max_row + 1
    for column, value in enumerate(values, 1):
        cell = write_cell(sheet, row, column, value)
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
    return row


def band(sheet, row, title):
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    cell = write_cell(sheet, row, 1, title)
    cell.fill = PatternFill("solid", fgColor="245548")
    cell.font = Font(name="Calibri", color="FFFFFF", bold=True, size=12)
    sheet.row_dimensions[row].height = 25


def interpretation(state):
    if state == "OK":
        return "Verificado", "Si", "Cumple la comprobacion implementada."
    if state.startswith("NO_APLICA"):
        return "Fuera de alcance", "No", "Excluido del porcentaje; no demuestra que este migrado."
    if state.startswith("NO_COMPARABLE"):
        return "No comparable", "Si", "La reescritura impide verificar igualdad por hashes."
    if state == "ADVERTENCIA":
        return "Pendiente / parcial", "Si", "Revisar las incidencias; no se cuenta como verificado."
    if state == "NO_EVALUADO":
        return "Sin verificar", "Si", "No se pudo completar esta comprobacion."
    return "No conforme", "Si", "Revisar diferencias o fallos registrados."


def build_report(data, output_path):
    results = list(data.get("resultados_por_repositorio", []))
    events = data.get("errores", []) + data.get("advertencias", [])
    existing = {result["repositorio"] for result in results}
    for event in events:
        repository = event.get("repositorio", "GLOBAL")
        if repository != "GLOBAL" and repository not in existing:
            results.append({"repositorio": repository, "resultado": "ERROR",
                            "validaciones": {"disponibilidad_para_certificar": "NO_EVALUADO"},
                            "metadatos": {}, "detalle": "Validacion interrumpida; consulte las incidencias."})
            existing.add(repository)
    if not results:
        raise ValueError("No hay repositorios para generar hojas de evidencia.")

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "Evidencia de migracion GitLab a GitHub"
    workbook.properties.creator = "Validador de migracion"
    metadata = data.get("metadatos", {})
    used_titles = set()
    for result in results:
        repository = str(result["repositorio"])
        sheet = workbook.create_sheet(sheet_title(repository, used_titles))
        sheet.sheet_view.showGridLines = False
        band(sheet, 1, "EVIDENCIA DE MIGRACION - " + repository)
        checks = list(result.get("validaciones", {}).items()) + list(result.get("metadatos", {}).items())
        applicable = sum(not str(state).startswith("NO_APLICA") for _, state in checks)
        verified = sum(state == "OK" for _, state in checks)
        percentage = math.floor(verified * 10000 / applicable + 0.5) / 10000 if applicable else None
        append_row(sheet, ["Repositorio", repository, "Estado", result.get("resultado", "NO_EVALUADO")])
        append_row(sheet, ["GitHub destino", metadata.get("organizacion_github", ""), "Grupo GitLab", metadata.get("grupo_gitlab", "")])
        append_row(sheet, ["Inicio", metadata.get("inicio", ""), "Fin", metadata.get("fin", "")])
        append_row(sheet, ["Porcentaje verificado", percentage if percentage is not None else "N/D",
                           "Comprobaciones OK", verified, "Aplicables", applicable])
        if percentage is not None:
            sheet["B5"].number_format = "0.00%"
            sheet.conditional_formatting.add("B5", DataBarRule(start_type="num", start_value=0,
                                                               end_type="num", end_value=1, color="2F8067"))
        sheet.merge_cells("A6:F6")
        write_cell(sheet, 6, 1, "Porcentaje de comprobaciones aplicables en OK, con igual peso. No mide bytes ni archivos; 0% sin verificar no significa repositorio vacio.")
        sheet.row_dimensions[6].height = 32
        sheet.merge_cells("A7:F7")
        write_cell(sheet, 7, 1, "NO_APLICA se excluye. Advertencias, errores y estados no comparables no suman. Solo se afirma el alcance implementado por el validador; GitHub no expone valores de Secrets.")
        sheet.row_dimensions[7].height = 42
        headers = ["Componente", "Estado tecnico", "Resultado", "Cuenta en porcentaje", "Comprobacion / limite", "Detalle"]
        header_row = append_row(sheet, headers, "DDEBE6")
        for cell in sheet[header_row]:
            cell.font = Font(bold=True, color="173C31")
        for component, state in checks:
            state = str(state)
            outcome, counted, detail = interpretation(state)
            append_row(sheet, [LABELS.get(component, component), state, outcome, counted, detail, ""], COLORS.get(state, "F1F3F4"))
        sheet.auto_filter.ref = f"A{header_row}:F{sheet.max_row}"
        sheet.freeze_panes = f"C{header_row + 1}"
        row = sheet.max_row + 2
        band(sheet, row, "INCIDENCIAS Y EVIDENCIA")
        append_row(sheet, ["Fecha", "Severidad", "Componente / etapa", "Mensaje", "Evidencia", "Continuacion"], "DDEBE6")
        incidents = [event for event in events if event.get("repositorio") in (repository, "GLOBAL")]
        if not incidents:
            append_row(sheet, ["", "", "", "Sin incidencias registradas en el alcance configurado."])
        for event in incidents:
            evidence = str(event.get("evidencia", ""))
            parts = [evidence[start:start + 30000] for start in range(0, len(evidence), 30000)] or [""]
            for index, part in enumerate(parts):
                row = append_row(sheet, [event.get("fecha", ""), event.get("severidad", ""),
                                          event.get("etapa", ""), event.get("mensaje", "") if index == 0 else "Continuacion de evidencia",
                                          part, f"{index + 1}/{len(parts)}"], COLORS.get(event.get("severidad")))
                sheet.row_dimensions[row].height = min(180, max(45, 15 * (part.count("\n") + 2)))
        row = sheet.max_row + 2
        band(sheet, row, "TRAZABILIDAD")
        append_row(sheet, ["Resultado", result.get("detalle", "")])
        append_row(sheet, ["GitLab API", metadata.get("gitlab_url", "")])
        append_row(sheet, ["Log de ejecucion", metadata.get("log_principal", "")])
        append_row(sheet, ["Generado", metadata.get("fecha_generacion", "")])
        for column, width in enumerate([37, 31, 32, 38, 65, 20], 1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = f"1:{header_row}"
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return len(results)


def main():
    parser = argparse.ArgumentParser(description="Exporta evidencia de migracion a Excel, una hoja por repositorio.")
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    arguments = parser.parse_args()
    with arguments.input.open(encoding="utf-8-sig") as source:
        data = json.load(source)
    count = build_report(data, arguments.output)
    print(f"Reporte Excel generado: {arguments.output} ({count} hojas)")


if __name__ == "__main__":
    main()